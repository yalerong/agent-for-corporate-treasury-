"""摄取管线 review 补丁：集合绑定、中断更新不留孤儿块、docx 表格顺序。"""
from __future__ import annotations

import pytest
from qdrant_client import QdrantClient

from app.rag import get_store
from app.rag import ingest as ingest_module
from app.rag.ingest import Manifest, parse_file, sync_directory
from tests.test_rag import DeterministicFakeEmbeddings


def _store(name: str = "t_review"):
    return get_store(
        "enterprise", DeterministicFakeEmbeddings(),
        client=QdrantClient(location=":memory:"), collection_name=name,
    )


@pytest.fixture
def corpus(tmp_dir):
    root = tmp_dir / "corpus"
    root.mkdir()
    (root / "transfer.md").write_text("调拨制度：跨主体须双签。", "utf-8")
    (root / "fx.txt").write_text("购汇额度不得超过预算。", "utf-8")
    return root


def _sync(root, store, tmp_dir, **kw):
    m = Manifest(tmp_dir / "m.json")
    return sync_directory(root, "enterprise", store, m, **kw), m


def _count(store) -> int:
    return store.client.count(store.collection_name, exact=True).count


def test_manifest_bound_to_its_collection(corpus, tmp_dir):
    _sync(corpus, _store("a"), tmp_dir)
    other = _store("b")
    r, _ = _sync(corpus, other, tmp_dir)
    assert r.blocked and "'b'" in r.blocked and _count(other) == 0


def test_manifest_with_empty_collection_is_blocked(corpus, tmp_dir):
    _sync(corpus, _store("a"), tmp_dir)
    r, _ = _sync(corpus, _store("a"), tmp_dir)  # 同名但是新建的空集合
    assert r.blocked and "集合是空的" in r.blocked


def test_interrupted_replace_leaves_no_orphans(corpus, tmp_dir, monkeypatch):
    store = _store()
    _sync(corpus, store, tmp_dir)
    (corpus / "fx.txt").write_text("购汇新规第一版。", "utf-8")

    def boom(*a, **k):
        raise RuntimeError("断了")

    monkeypatch.setattr(ingest_module, "_delete_source", boom)
    with pytest.raises(RuntimeError):
        _sync(corpus, store, tmp_dir)  # 新块已写入、旧块没删、manifest 没更新
    monkeypatch.undo()

    (corpus / "fx.txt").unlink()  # 随后这份文档被删除
    r, m = _sync(corpus, store, tmp_dir)
    assert r.deleted == ["fx.txt"] and "fx.txt" not in m.docs
    texts = [h.page_content for h in store.similarity_search("购汇", k=20)]
    assert not any("购汇" in t for t in texts)  # 中断留下的新块也一并清掉了


def test_changed_file_cleans_chunks_missing_from_manifest(corpus, tmp_dir):
    """中断后再次编辑：旧账里没有的中间版本块，在下一次成功入库时被清掉。"""
    store = _store()
    _sync(corpus, store, tmp_dir)
    (corpus / "fx.txt").write_text("购汇第二版。", "utf-8")
    _sync(corpus, store, tmp_dir)
    # 模拟一批没记账的残留块（同 source、别的哈希）
    store.add_texts(
        ["购汇残留块。"], metadatas=[{"source": "fx.txt", "doc_sha": "stale"}]
    )
    (corpus / "fx.txt").write_text("购汇第三版。", "utf-8")
    _sync(corpus, store, tmp_dir)
    texts = sorted(h.page_content for h in store.similarity_search("购汇", k=20) if "购汇" in h.page_content)
    assert texts == ["购汇第三版。"]


def test_docx_keeps_paragraph_table_order(tmp_dir):
    from docx import Document

    d = Document()
    d.add_heading("境内调拨", level=2)
    t1 = d.add_table(rows=1, cols=2)
    t1.rows[0].cells[0].text = "境内限额"
    t1.rows[0].cells[1].text = "50万"
    d.add_heading("跨境调拨", level=2)
    t2 = d.add_table(rows=1, cols=2)
    t2.rows[0].cells[0].text = "跨境限额"
    t2.rows[0].cells[1].text = "10万"
    p = tmp_dir / "o.docx"
    d.save(str(p))
    text = parse_file(p)
    order = [text.index(x) for x in ("境内调拨", "境内限额", "跨境调拨", "跨境限额")]
    assert order == sorted(order)
