"""知识库摄取管线：解析、增量同步、删除对账、失败保旧、批量删除保险丝。"""
from __future__ import annotations

import json
from pathlib import Path

import pytest
from langchain_core.documents import Document
from qdrant_client import QdrantClient

from app.rag import get_store
from app.rag.ingest import (
    Manifest,
    ParseError,
    parse_file,
    scan_root,
    sync_directory,
)
from tests.test_rag import DeterministicFakeEmbeddings


@pytest.fixture
def store():
    client = QdrantClient(location=":memory:")
    return get_store(
        "enterprise", DeterministicFakeEmbeddings(), client=client, collection_name="t_ingest"
    )


@pytest.fixture
def corpus(tmp_dir):
    root = tmp_dir / "corpus"
    (root / "policies").mkdir(parents=True)
    (root / "policies" / "transfer.md").write_text("# 调拨制度\n跨主体调拨须双签。", "utf-8")
    (root / "policies" / "fx.txt").write_text("购汇额度不得超过预算。", "utf-8")
    return root


def _count(store) -> int:
    return store.client.count(store.collection_name, exact=True).count


def _sync(corpus, store, manifest_dir, **kw):
    m = Manifest(manifest_dir / "m.json")
    return sync_directory(corpus, "enterprise", store, m, **kw), m


class TestParse:
    def test_gbk_text(self, tmp_dir):
        p = tmp_dir / "a.txt"
        p.write_bytes("资金调拨".encode("gb18030"))
        assert parse_file(p) == "资金调拨"

    def test_empty_file_is_error(self, tmp_dir):
        p = tmp_dir / "a.md"
        p.write_text("  \n", "utf-8")
        with pytest.raises(ParseError):
            parse_file(p)

    def test_unsupported_suffix(self, tmp_dir):
        p = tmp_dir / "a.png"
        p.write_bytes(b"x")
        with pytest.raises(ParseError):
            parse_file(p)

    def test_docx_headings_and_tables(self, tmp_dir):
        from docx import Document

        d = Document()
        d.add_heading("付款制度", level=1)
        d.add_paragraph("大额付款须复核。")
        t = d.add_table(rows=1, cols=2)
        t.rows[0].cells[0].text = "金额"
        t.rows[0].cells[1].text = "审批人"
        p = tmp_dir / "a.docx"
        d.save(str(p))
        text = parse_file(p)
        assert "# 付款制度" in text and "大额付款须复核。" in text and "金额 | 审批人" in text

    def test_xlsx_sheet_rows(self, tmp_dir):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "阈值"
        ws.append(["档位", "金额"])
        ws.append(["一级", 50000])
        p = tmp_dir / "a.xlsx"
        wb.save(str(p))
        text = parse_file(p)
        assert "## 阈值" in text and "一级 | 50000" in text

    def test_pdf_without_text_layer_is_error(self, tmp_dir):
        from pypdf import PdfWriter

        w = PdfWriter()
        w.add_blank_page(width=100, height=100)
        p = tmp_dir / "scan.pdf"
        with p.open("wb") as f:
            w.write(f)
        with pytest.raises(ParseError):
            parse_file(p)

    def test_corrupt_docx_is_parse_error(self, tmp_dir):
        p = tmp_dir / "bad.docx"
        p.write_bytes(b"not a zip")
        with pytest.raises(ParseError):
            parse_file(p)


def test_scan_skips_temp_hidden_and_unsupported(corpus):
    (corpus / "~$lock.docx").write_bytes(b"x")
    (corpus / ".git").mkdir()
    (corpus / ".git" / "n.md").write_text("x", "utf-8")
    (corpus / "pic.png").write_bytes(b"x")
    assert sorted(scan_root(corpus)) == ["policies/fx.txt", "policies/transfer.md"]


class TestSync:
    def test_first_run_adds_all_and_rerun_is_noop(self, corpus, store, tmp_dir):
        r1, _ = _sync(corpus, store, tmp_dir)
        assert sorted(r1.added) == ["policies/fx.txt", "policies/transfer.md"]
        n = _count(store)
        assert n >= 2
        r2, _ = _sync(corpus, store, tmp_dir)
        assert not r2.added and not r2.changed and len(r2.unchanged) == 2
        assert _count(store) == n

    def test_changed_file_replaces_old_chunks(self, corpus, store, tmp_dir):
        _sync(corpus, store, tmp_dir)
        (corpus / "policies" / "fx.txt").write_text("购汇额度上限调整为预算的 80%。", "utf-8")
        r, m = _sync(corpus, store, tmp_dir)
        assert r.changed == ["policies/fx.txt"]
        hits = store.similarity_search("购汇额度", k=10)
        texts = [h.page_content for h in hits if h.metadata["source"] == "policies/fx.txt"]
        assert texts == ["购汇额度上限调整为预算的 80%。"]
        assert _count(store) == sum(len(d["chunk_ids"]) for d in m.docs.values())

    def test_deleted_file_removes_chunks(self, corpus, store, tmp_dir):
        # 三份里删一份，未超过保险丝比例
        (corpus / "policies" / "c.md").write_text("第三份。", "utf-8")
        _sync(corpus, store, tmp_dir)
        (corpus / "policies" / "c.md").unlink()
        r, m = _sync(corpus, store, tmp_dir)
        assert r.deleted == ["policies/c.md"]
        assert "policies/c.md" not in m.docs
        assert _count(store) == sum(len(d["chunk_ids"]) for d in m.docs.values())

    def test_parse_failure_keeps_old_chunks(self, corpus, store, tmp_dir):
        _sync(corpus, store, tmp_dir)
        n = _count(store)
        (corpus / "policies" / "fx.txt").write_text("   ", "utf-8")
        r, m = _sync(corpus, store, tmp_dir)
        assert "policies/fx.txt" in r.errors
        assert "policies/fx.txt" not in r.changed
        assert _count(store) == n  # 旧块还在
        assert m.docs["policies/fx.txt"]["sha256"]  # manifest 仍是旧版本，下次会重试

    def test_new_file_parse_failure_not_indexed(self, corpus, store, tmp_dir):
        (corpus / "empty.md").write_text("", "utf-8")
        r, m = _sync(corpus, store, tmp_dir)
        assert "empty.md" in r.errors and "empty.md" not in r.added
        assert "empty.md" not in m.docs

    def test_mass_delete_blocked_unless_forced(self, corpus, store, tmp_dir):
        _sync(corpus, store, tmp_dir)
        n = _count(store)
        for f in corpus.rglob("*.*"):
            f.unlink()
        r, m = _sync(corpus, store, tmp_dir)
        assert r.blocked and not r.deleted
        assert _count(store) == n and len(m.docs) == 2
        r2, m2 = _sync(corpus, store, tmp_dir, force=True)
        assert len(r2.deleted) == 2 and _count(store) == 0 and not m2.docs

    def test_missing_root_is_blocked_not_wiped(self, corpus, store, tmp_dir):
        _sync(corpus, store, tmp_dir)
        n = _count(store)
        r, _ = _sync(tmp_dir / "gone", store, tmp_dir)
        assert r.blocked and _count(store) == n

    def test_dry_run_writes_nothing(self, corpus, store, tmp_dir):
        r, m = _sync(corpus, store, tmp_dir, dry_run=True)
        assert len(r.added) == 2 and _count(store) == 0 and not m.docs
        assert not (tmp_dir / "m.json").exists()

    def test_dry_run_needs_no_store(self, corpus, tmp_dir):
        # 预演只读文件系统与 manifest：不加载模型、不建集合
        r = sync_directory(
            corpus, "enterprise", None, Manifest(tmp_dir / "m.json"), dry_run=True
        )
        assert sorted(r.added) == ["policies/fx.txt", "policies/transfer.md"]
        assert not (tmp_dir / "m.json").exists()

    def test_unreadable_directory_skips_deletion(self, corpus, store, tmp_dir, monkeypatch):
        _sync(corpus, store, tmp_dir)
        n = _count(store)
        blocked_dir = corpus / "policies"
        real_iterdir = Path.iterdir

        def fake_iterdir(self):
            if self == blocked_dir:
                raise PermissionError(13, "Access denied")
            return real_iterdir(self)

        monkeypatch.setattr(Path, "iterdir", fake_iterdir)
        r, m = _sync(corpus, store, tmp_dir)
        assert r.scan_errors and not r.deleted
        assert _count(store) == n and len(m.docs) == 2

    def test_unreadable_file_is_per_document_error(self, corpus, store, tmp_dir, monkeypatch):
        _sync(corpus, store, tmp_dir)
        n = _count(store)
        (corpus / "policies" / "fx.txt").write_text("购汇额度上限调整为预算的 80%。", "utf-8")
        from app.rag import ingest as ingest_module

        real_sha = ingest_module.file_sha256

        def fake_sha(path):
            if path.name == "fx.txt":
                raise PermissionError(13, "Access denied")
            return real_sha(path)

        monkeypatch.setattr(ingest_module, "file_sha256", fake_sha)
        r, m = _sync(corpus, store, tmp_dir)
        assert "读取失败" in r.errors["policies/fx.txt"]
        assert _count(store) == n and m.docs["policies/fx.txt"]["sha256"]  # 旧版本原样保留

    def test_legacy_collection_requires_forced_rebuild(self, corpus, store, tmp_dir):
        # 旧索引：随机块 ID、没有 manifest。对不了账，继续同步只会叠重复块
        store.add_documents([Document(page_content="旧索引", metadata={"source": "policies/fx.txt"})])
        r, m = _sync(corpus, store, tmp_dir)
        assert r.blocked and "manifest" in r.blocked and not m.docs
        r2, m2 = _sync(corpus, store, tmp_dir, force=True)
        assert not r2.blocked and r2.warnings
        assert sorted(m2.docs) == ["policies/fx.txt", "policies/transfer.md"]
        assert _count(store) == sum(len(d["chunk_ids"]) for d in m2.docs.values())

    def test_legacy_manifest_entries_are_upgraded(self, corpus, store, tmp_dir):
        _sync(corpus, store, tmp_dir)
        n = _count(store)
        path = tmp_dir / "m.json"
        raw = json.loads(path.read_text("utf-8"))
        for entry in raw["docs"].values():  # 上一版 manifest：没有 meta / superseded_on
            entry.pop("meta")
            entry.pop("superseded_on")
        path.write_text(json.dumps(raw, ensure_ascii=False), "utf-8")

        # dry-run 也要能读旧格式，不能 KeyError
        dry, _ = _sync(corpus, store, tmp_dir, dry_run=True)
        assert not dry.errors and sorted(dry.changed) == [
            "policies/fx.txt", "policies/transfer.md"
        ]

        r, m = _sync(corpus, store, tmp_dir)
        assert not r.errors and any("旧格式" in w for w in r.warnings)
        assert all("meta" in d for d in m.docs.values())
        assert _count(store) == n  # 旧块 ID 照样对齐、不产生重复

    def test_source_prefix_migration_keeps_existing_docs(self, corpus, store, tmp_dir):
        # 老 manifest 的 source 不带 track 前缀；换成带前缀的命名空间时不能当成"文档全没了"
        _sync(corpus, store, tmp_dir)
        n = _count(store)
        r, m = _sync(corpus, store, tmp_dir, source_prefix="enterprise/")
        assert not r.blocked and not r.deleted and not r.errors
        assert sorted(m.docs) == [
            "enterprise/policies/fx.txt", "enterprise/policies/transfer.md"
        ]
        assert any("前缀" in w for w in r.warnings)
        assert _count(store) == n  # 旧块按 source 重建后删掉，不残留也不重复
        sources = {h.metadata["source"] for h in store.similarity_search("调拨", k=10)}
        assert sources == {"enterprise/policies/fx.txt", "enterprise/policies/transfer.md"}

    def test_chunk_metadata_and_stable_ids(self, corpus, store, tmp_dir):
        _, m1 = _sync(corpus, store, tmp_dir)
        hit = store.similarity_search("跨主体调拨", k=1)[0]
        assert hit.metadata["source"] == "policies/transfer.md"
        assert hit.metadata["doc_sha"] and "chunk_index" in hit.metadata
        # 清空 manifest 重跑：集合里已有块但没法对账 → 先拒绝
        n = _count(store)
        (tmp_dir / "m.json").unlink()
        r, _ = _sync(corpus, store, tmp_dir)
        assert r.blocked and _count(store) == n
        # --force 重建：确定性 ID 不变 → 块数不变，不产生重复
        r2, _ = _sync(corpus, store, tmp_dir, force=True)
        assert not r2.blocked and _count(store) == n

    def test_empty_collection_with_manifest_is_blocked(self, corpus, store, tmp_dir):
        # 换向量库后端 / 数据卷丢了：manifest 还在但集合是空的，不能报"未变"
        _sync(corpus, store, tmp_dir)
        manifest = Manifest(tmp_dir / "m.json")
        store.delete(ids=[i for d in manifest.docs.values() for i in d["chunk_ids"]])
        assert _count(store) == 0
        r, _ = _sync(corpus, store, tmp_dir)
        assert r.blocked and not r.unchanged and "集合是空的" in r.blocked

    def test_scan_error_keeps_id_reservations(self, corpus, store, tmp_dir, monkeypatch):
        _sync(corpus, store, tmp_dir)
        # 可读目录里放一份冒用不可读文档 doc_id 的新文档
        (corpus / "copy.md").write_text("---\ndoc_id: policies/fx.txt\n---\n冒用。", "utf-8")
        blocked_dir = corpus / "policies"
        real_iterdir = Path.iterdir

        def fake_iterdir(self):
            if self == blocked_dir:
                raise PermissionError(13, "Access denied")
            return real_iterdir(self)

        monkeypatch.setattr(Path, "iterdir", fake_iterdir)
        r, m = _sync(corpus, store, tmp_dir)
        assert r.scan_errors and "已被" in r.errors["copy.md"]
        ids = [d["meta"]["doc_id"] for d in m.docs.values()]
        assert len(ids) == len(set(ids))  # 不可读文档的 id 没被抢走

    def test_dry_run_checks_proposed_metadata(self, corpus, store, tmp_dir):
        # 预演也要按"拟入库"的元数据出草稿清单和取代关系告警，且不污染调用方的 manifest
        (corpus / "draft.md").write_text("无元数据。", "utf-8")
        (corpus / "new.md").write_text(
            "---\nstatus: approved\ndoc_id: n2\nsupersedes: [ghost]\n---\n新制度。", "utf-8"
        )
        r, m = _sync(corpus, store, tmp_dir, dry_run=True)
        assert "draft.md" in r.drafts
        assert any("ghost" in w for w in r.warnings)
        assert not m.docs and _count(store) == 0

    def test_dry_run_drops_proposed_deletions_before_validation(self, corpus, store, tmp_dir):
        # 已入库三份：old(approved, id=o1) / dr(草稿) / keep；之后 old、dr 被删，新增一份取代 o1
        (corpus / "old.md").write_text("---\nstatus: approved\ndoc_id: o1\n---\n旧。", "utf-8")
        (corpus / "dr.md").write_text("草稿，无元数据。", "utf-8")
        (corpus / "keep.md").write_text("---\nstatus: approved\n---\n保留。", "utf-8")
        _sync(corpus, store, tmp_dir)
        (corpus / "old.md").unlink()
        (corpus / "dr.md").unlink()
        (corpus / "new.md").write_text(
            "---\nstatus: approved\ndoc_id: n1\nsupersedes: [o1]\n---\n新。", "utf-8"
        )
        n = _count(store)
        r, m = _sync(corpus, store, tmp_dir, dry_run=True)
        assert sorted(r.deleted) == ["dr.md", "old.md"]
        # o1 在同步后已不存在 → 与真实同步一样报“知识库里没有这份文档”
        assert any("o1" in w for w in r.warnings)
        assert "dr.md" not in r.drafts                      # 已删的草稿不再挂在待批准清单上
        assert {"old.md", "dr.md"} <= set(m.docs) and _count(store) == n  # 调用方的 manifest 与库都没动


def test_index_track_propagates_failure(monkeypatch, tmp_dir):
    """批量入口不能把"某份文档失败/整批被拒"当成刷新成功。"""
    from app.rag import store as store_module
    from app.rag.ingest import IngestError

    bad = tmp_dir / "kb"
    (bad / "policies").mkdir(parents=True)
    (bad / "policies" / "x.md").write_text("---\nstaus: approved\n---\n正文。", "utf-8")
    fake_store = get_store(
        "enterprise", DeterministicFakeEmbeddings(),
        client=QdrantClient(location=":memory:"), collection_name="t_track",
    )
    monkeypatch.setattr(store_module, "get_embeddings", lambda: DeterministicFakeEmbeddings())
    monkeypatch.setattr(store_module, "get_store", lambda track, embeddings: fake_store)
    settings = store_module.get_settings()
    monkeypatch.setattr(settings, "enterprise_kb_path", bad)
    monkeypatch.setattr(settings, "qdrant_path", tmp_dir / "qd")

    with pytest.raises(IngestError):
        store_module.index_track("enterprise")
