"""企业知识库摄取管线：多格式解析 + 内容哈希增量同步。

和 `index_track` 的区别：后者是"全量重建、只认 .md/.txt"，这里是生产路径——
- 解析 .md/.txt/.docx/.pdf/.xlsx/.csv（重依赖延迟导入，纯文本场景不装也能跑）
- 每份文档记哈希，下次只处理 新增 / 改动 / 删除 三类；源头删了，索引里的块也跟着删
- 块 ID 由 (source, 内容哈希, 序号) 确定性生成，重跑不产生重复块
- 单份文档解析失败：保留它旧的块不动（宁可旧，不可空），错误写进报告
- 一次要删掉超过一半的已入库文档 → 默认拒绝（挂载点掉了/路径写错的典型症状），--force 才放行

source 是相对 root 的路径（不含 root 目录名），换机器、换挂载点不变。
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import sys
import uuid
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path

from langchain_core.documents import Document

from app.rag.docmeta import (
    NEVER,
    DocMeta,
    MetaError,
    compute_superseded_on,
    inputs_digest,
    resolve_meta,
    split_front_matter,
)
from app.rag.store import Track, chunk_documents

SUPPORTED_SUFFIXES = (".md", ".txt", ".docx", ".pdf", ".xlsx", ".csv")
MASS_DELETE_RATIO = 0.5
_NS = uuid.UUID("6f1c2f0e-3b8e-4c43-9d57-5a0c1b7d9e11")


class ParseError(Exception):
    """文件读不出可用文本（损坏、加密、扫描件无文字层等）。"""


class IngestError(RuntimeError):
    """整批同步没有干净完成（被拒绝执行或有文档失败），调用方不能当成功处理。"""


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "gb18030"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ParseError("无法按 utf-8 / gb18030 解码")


def _parse_docx(path: Path) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    parts: list[str] = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "") if p.style is not None else ""
        if style.startswith("Heading"):
            level = "".join(c for c in style if c.isdigit()) or "1"
            parts.append("#" * min(int(level), 6) + " " + text)
        else:
            parts.append(text)
    for t in doc.tables:
        rows = [" | ".join(c.text.strip().replace("\n", " ") for c in r.cells) for r in t.rows]
        parts.append("\n".join(rows))
    return "\n\n".join(parts)


def _parse_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        raise ParseError("PDF 已加密")
    parts = []
    for i, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if text:
            parts.append(f"[第{i}页]\n{text}")
    return "\n\n".join(parts)


def _parse_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    try:
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(values_only=True):
                cells = ["" if v is None else str(v).strip() for v in row]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                parts.append(f"## {ws.title}\n" + "\n".join(rows))
    finally:
        wb.close()
    return "\n\n".join(parts)


_PARSERS: dict[str, Callable[[Path], str]] = {
    ".md": _read_text,
    ".txt": _read_text,
    ".csv": _read_text,
    ".docx": _parse_docx,
    ".pdf": _parse_pdf,
    ".xlsx": _parse_xlsx,
}


def parse_file(path: Path) -> str:
    """按后缀解析成纯文本。读不出任何文字 → ParseError（扫描件需 OCR，本管线不做）。"""
    parser = _PARSERS.get(path.suffix.lower())
    if parser is None:
        raise ParseError(f"不支持的格式: {path.suffix}")
    try:
        text = parser(path)
    except ParseError:
        raise
    except Exception as e:  # 第三方库对损坏文件抛的异常类型各异
        raise ParseError(f"{type(e).__name__}: {e}") from e
    if not text.strip():
        raise ParseError("解析后无文字（扫描件或空文件）")
    return text


def file_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


class Manifest:
    """source → {sha256, chunk_ids, indexed_at}。每处理完一份文档就落盘（原子替换）。"""

    def __init__(self, path: Path):
        self.path = path
        self.docs: dict[str, dict] = {}
        if path.exists():
            self.docs = json.loads(path.read_text(encoding="utf-8")).get("docs", {})

    def save(self) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        tmp = self.path.with_suffix(self.path.suffix + ".tmp")
        tmp.write_text(
            json.dumps({"version": 1, "docs": self.docs}, ensure_ascii=False, indent=1),
            encoding="utf-8",
        )
        os.replace(tmp, self.path)


@dataclass
class SyncReport:
    added: list[str] = field(default_factory=list)
    changed: list[str] = field(default_factory=list)
    unchanged: list[str] = field(default_factory=list)
    deleted: list[str] = field(default_factory=list)
    errors: dict[str, str] = field(default_factory=dict)
    scan_errors: list[str] = field(default_factory=list)
    drafts: list[str] = field(default_factory=list)  # 已入库但非 approved：检索不到
    warnings: list[str] = field(default_factory=list)
    blocked: str | None = None
    dry_run: bool = False

    def summary(self) -> str:
        head = "[dry-run] " if self.dry_run else ""
        s = (
            f"{head}新增 {len(self.added)}｜改动 {len(self.changed)}｜"
            f"删除 {len(self.deleted)}｜未变 {len(self.unchanged)}｜失败 {len(self.errors)}"
        )
        if self.drafts:
            s += (
                f"\n其中 {len(self.drafts)} 份尚未批准（status 不是 approved），检索不到；"
                "在文档 front matter / .meta.yaml / _meta.yaml 里写 status: approved 后重跑"
            )
        for w in self.warnings:
            s += f"\n提示: {w}"
        if self.scan_errors:
            s += (
                f"\n扫描不完整（{len(self.scan_errors)} 个目录读不了），"
                "已跳过删除对账: " + "；".join(self.scan_errors)
            )
        if self.blocked:
            s += f"\n已拒绝执行: {self.blocked}"
        return s


def scan_root(
    root: Path, prefix: str = "", errors: list[str] | None = None
) -> dict[str, Path]:
    """root 下所有受支持文件 → {prefix+相对路径(posix): 绝对路径}。跳过 ~$ 临时文件与隐藏条目。

    传了 errors 时，读不了的目录记进 errors 并继续（调用方据此拒绝做删除对账）；
    不传则沿用旧行为：目录读不了直接抛 OSError。
    """
    found: dict[str, Path] = {}

    def walk(directory: Path) -> None:
        try:
            entries = sorted(directory.iterdir())
        except OSError as e:
            if errors is None:
                raise
            errors.append(f"{directory}: {e.strerror or e}")
            return
        for entry in entries:
            if entry.name.startswith("."):
                continue
            if entry.is_dir():
                # 符号链接目录不跟，避免环
                if not entry.is_symlink():
                    walk(entry)
            elif not entry.name.startswith("~$") and entry.suffix.lower() in SUPPORTED_SUFFIXES:
                found[prefix + entry.relative_to(root).as_posix()] = entry

    walk(root)
    return found


def _chunk_ids(source: str, sha: str, n: int) -> list[str]:
    return [str(uuid.uuid5(_NS, f"{source}#{sha}#{i}")) for i in range(n)]


def _drop_all(store) -> int:
    """清空集合里所有块（显式重建用：没有 manifest 时无法逐块对账）。"""
    deleted = 0
    while True:
        points, _ = store.client.scroll(
            store.collection_name, limit=256, with_payload=False, with_vectors=False
        )
        if not points:
            return deleted
        store.delete(ids=[str(p.id) for p in points])
        deleted += len(points)


def _migrate_source_prefix(
    manifest: Manifest,
    on_disk: dict[str, Path],
    prefix: str,
    report: SyncReport,
    dry_run: bool,
) -> None:
    """旧 manifest 的 source 不带 `<track>/` 前缀时，按前缀迁一次，而不是当成文档消失。

    迁移后 sha256 清空 → 本轮按新 source 重建块（块 ID 含 source），旧块 ID 先留着，
    重建成功后按 stale 删掉；引用它的 doc_id 仍指向同一份文档，不会重复占坑。
    """
    if not prefix:
        return
    moved = 0
    for source in [s for s in manifest.docs if s not in on_disk]:
        new_source = prefix + source
        if new_source not in on_disk or new_source in manifest.docs:
            continue
        entry = manifest.docs.pop(source)
        entry["sha256"] = ""
        manifest.docs[new_source] = entry
        moved += 1
    if moved:
        report.warnings.append(
            f"manifest 里 {moved} 份文档的 source 补上 {prefix} 前缀，本轮重新入库"
        )
        if not dry_run:
            manifest.save()


def _upgrade_legacy_entries(
    manifest: Manifest, report: SyncReport, dry_run: bool
) -> None:
    """把上一版 manifest（只有 sha256/chunk_ids/indexed_at）补成带 meta 的格式。

    补的 meta 只是占位：sha256 清空 → 本轮回读元数据并按新格式重建块，
    旧块 ID 保留在 entry 里，重建成功后按 stale 删除。
    """
    legacy = [src for src, d in manifest.docs.items() if "meta" not in d]
    for src in legacy:
        entry = manifest.docs[src]
        entry["meta"] = DocMeta(doc_id=src).to_dict()
        entry["sha256"] = ""
        entry.setdefault("superseded_on", NEVER)
    if legacy:
        report.warnings.append(
            f"manifest 里 {len(legacy)} 份旧格式条目缺元数据，已标记重新入库"
        )
        if not dry_run:
            manifest.save()


def _load(root: Path, path: Path, source: str) -> tuple[DocMeta, str]:
    """解析正文 + 元数据。文本类格式的 front matter 在这里拆掉，不进索引。"""
    text = parse_file(path)
    front = None
    if path.suffix.lower() in (".md", ".txt"):
        front, text = split_front_matter(text)
        if not text.strip():
            raise ParseError("去掉 front matter 后正文为空")
    return resolve_meta(root, path, source, front), text


def _refresh_superseded(store, manifest: Manifest, report: SyncReport, dry_run: bool) -> None:
    """按全库已批准文档重算 superseded_on，只对变化的文档改块 payload。"""
    computed, warnings = compute_superseded_on(
        {src: d["meta"] for src, d in manifest.docs.items()}
    )
    report.warnings.extend(warnings)
    for src, value in computed.items():
        d = manifest.docs[src]
        if d.get("superseded_on", NEVER) == value or dry_run:
            continue
        if d["chunk_ids"]:
            store.client.set_payload(
                store.collection_name,
                payload={"superseded_on": value},
                points=d["chunk_ids"],
                key=store.METADATA_KEY,
            )
        d["superseded_on"] = value
        manifest.save()


def sync_directory(
    root: Path,
    category: str,
    store,
    manifest: Manifest,
    *,
    dry_run: bool = False,
    force: bool = False,
    chunk_size: int = 500,
    chunk_overlap: int = 50,
    source_prefix: str = "",
) -> SyncReport:
    """把 root 与 manifest 对账，增量写入 store。返回逐文档的处理报告。

    dry_run=True 时只读文件系统与 manifest，不碰 store（可以传 store=None）。
    """
    report = SyncReport(dry_run=dry_run)
    if not root.is_dir():
        report.blocked = f"root 不存在或不是目录: {root}"
        return report
    if store is None and not dry_run:
        raise ValueError("store 不能为空（只有 dry_run 可以不带）")

    on_disk = scan_root(root, source_prefix, errors=report.scan_errors)
    _migrate_source_prefix(manifest, on_disk, source_prefix, report, dry_run)
    _upgrade_legacy_entries(manifest, report, dry_run)
    # 目录读不全时，"文件没了" 可能只是看不见：宁可不对账，也绝不删错
    vanished = [] if report.scan_errors else sorted(set(manifest.docs) - set(on_disk))
    if (
        not report.scan_errors
        and manifest.docs
        and len(vanished) > len(manifest.docs) * MASS_DELETE_RATIO
        and not force
    ):
        report.blocked = (
            f"将删除 {len(vanished)}/{len(manifest.docs)} 份已入库文档，"
            "疑似挂载点丢失或路径错误；确认无误请加 --force"
        )
        return report

    if not dry_run and not manifest.docs:
        # manifest 为空但集合非空：多半是旧索引（scripts.build_kb 的随机块 ID）。
        # 没法逐块对账，继续同步只会叠加重复块，所以要求显式重建。
        existing = store.client.count(store.collection_name, exact=True).count
        if existing and not force:
            report.blocked = (
                f"集合里有 {existing} 个块，但没有 manifest（旧索引或 manifest 丢失）；"
                "继续会叠加重复块。确认重建请加 --force（先清空该集合再全量入库）"
            )
            return report
        if existing:
            removed = _drop_all(store)
            report.warnings.append(f"已清空 {removed} 个无 manifest 可对账的旧块，按源文件重建")
    elif not dry_run:
        # manifest 说有已入库文档，但集合里一个块都没有：换过向量库后端或数据卷丢了。
        # 别让"未变"把空库粉饰成正常，直接按源文件全量重建。
        existing = store.client.count(store.collection_name, exact=True).count
        if existing == 0:
            report.warnings.append(
                f"集合里一个块都没有，但 manifest 记着 {len(manifest.docs)} 份已入库文档"
                "（换过向量库后端或数据卷丢了？），本轮按源文件全量重建"
            )
            manifest.docs.clear()

    claimed = {d["meta"]["doc_id"]: src for src, d in manifest.docs.items() if src in on_disk}
    for source, path in on_disk.items():
        old = manifest.docs.get(source)
        try:
            digest = inputs_digest(root, path, file_sha256(path))
        except OSError as e:
            # 单份读不了不影响其余文档；旧块与 manifest 原样保留，下次再试
            report.errors[source] = f"读取失败: {e.strerror or e}"
            continue
        if old and old["sha256"] == digest:
            report.unchanged.append(source)
            continue
        try:
            meta, text = _load(root, path, source)
            owner = claimed.get(meta.doc_id)
            if owner is not None and owner != source:
                raise MetaError(f"doc_id {meta.doc_id!r} 已被 {owner} 使用")
        except (ParseError, MetaError) as e:
            # 读不出或元数据非法：旧块原样保留（宁可旧，不可空/不可错放行），新文件不入库
            report.errors[source] = str(e)
            continue
        if old and old["meta"]["doc_id"] != meta.doc_id:
            # 换了 doc_id：旧 id 的占坑要释放，否则同批次里复用该 id 的文档会被误拒
            claimed.pop(old["meta"]["doc_id"], None)
        claimed[meta.doc_id] = source
        (report.changed if old else report.added).append(source)
        if dry_run:
            continue
        base = Document(
            page_content=text,
            metadata={
                "source": source,
                "category": category,
                "subcategory": path.parent.name,
                "doc_sha": digest,
                "superseded_on": NEVER,
                **meta.payload(),
            },
        )
        chunks = chunk_documents([base], chunk_size=chunk_size, chunk_overlap=chunk_overlap)
        ids = _chunk_ids(source, digest, len(chunks))
        now = datetime.now(UTC).isoformat(timespec="seconds")
        for i, c in enumerate(chunks):
            c.metadata["chunk_index"] = i
            c.metadata["indexed_at"] = now
        store.add_documents(chunks, ids=ids)  # 先写新块，再删旧块：中途失败只会多不会缺
        if old:
            stale = [i for i in old["chunk_ids"] if i not in set(ids)]
            if stale:
                store.delete(ids=stale)
        manifest.docs[source] = {
            "sha256": digest, "chunk_ids": ids, "indexed_at": now,
            "meta": meta.to_dict(), "superseded_on": NEVER,
        }
        manifest.save()

    for source in vanished:
        report.deleted.append(source)
        if dry_run:
            continue
        ids = manifest.docs[source]["chunk_ids"]
        if ids:
            store.delete(ids=ids)
        del manifest.docs[source]
        manifest.save()

    _refresh_superseded(store, manifest, report, dry_run)
    report.drafts = sorted(
        src for src, d in manifest.docs.items() if d["meta"]["status"] != "approved"
    )
    return report


def main(argv: list[str] | None = None) -> int:
    from app.config import get_settings
    from app.rag.store import get_embeddings, get_store

    s = get_settings()
    p = argparse.ArgumentParser(description="增量同步一个目录到知识库")
    p.add_argument("--track", choices=["industry", "enterprise"], required=True)
    p.add_argument("--root", type=Path, help="源目录；缺省用该 track 的配置路径")
    p.add_argument("--source-prefix", help="加在每个 source 前的前缀；缺省用配置路径时为 <track>/")
    p.add_argument("--dry-run", action="store_true", help="只对账不写入（不加载模型、不碰向量库）")
    p.add_argument("--force", action="store_true", help="放行大批量删除，或重建无 manifest 的旧索引")
    args = p.parse_args(argv)

    track: Track = args.track
    root = args.root or (s.industry_kb_path if track == "industry" else s.enterprise_kb_path)
    manifest = Manifest(s.qdrant_path / f"manifest-{track}.json")
    store = None if args.dry_run else get_store(track, get_embeddings())
    prefix = args.source_prefix if args.source_prefix is not None else ("" if args.root else f"{track}/")
    report = sync_directory(
        root, track, store, manifest, dry_run=args.dry_run, force=args.force, source_prefix=prefix
    )
    print(report.summary())
    for src, err in report.errors.items():
        print(f"  ✗ {src}: {err}", file=sys.stderr)
    for src in report.drafts:
        print(f"  · 待批准: {src}")
    return 1 if report.blocked or report.errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
